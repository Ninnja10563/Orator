"""Generate independent OOXML fixtures or verify Orator's exported archive.

Test tooling only. Requires python-pptx==1.0.2; not an application dependency.
"""
from pathlib import Path
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image
from io import BytesIO


def generate():
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333333), Inches(7.5)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(0.75), Inches(10), Inches(1))
    run = box.text_frame.paragraphs[0].add_run()
    run.text = 'Independent fixture — A & B < C'
    run.font.size = Pt(36)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x17, 0x1C, 0x24)
    slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1), Inches(2), Inches(3), Inches(2))
    table = slide.shapes.add_table(2, 2, Inches(5), Inches(2), Inches(5), Inches(2)).table
    for row, values in enumerate([['Region', 'Revenue'], ['North', '42']]):
        for col, value in enumerate(values):
            table.cell(row, col).text = value
    data = BytesIO()
    Image.new('RGB', (16, 16), (34, 91, 191)).save(data, format='PNG')
    data.seek(0)
    slide.shapes.add_picture(data, Inches(10.5), Inches(5), Inches(1), Inches(1))
    slide.notes_slide.notes_text_frame.text = 'Independent speaker notes'
    second = deck.slides.add_slide(deck.slide_layouts[6])
    second.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2)).text = 'Second slide'
    deck.save('Tests/PresentationCoreTests/Fixtures/independent.pptx')


def verify(path):
    deck = Presentation(path)
    assert len(deck.slides) >= 1
    texts = [s.text for s in deck.slides[0].shapes if s.has_text_frame]
    assert any('Your next great idea' in text for text in texts), texts
    assert deck.slide_width == 1280 * 9525
    assert any(s.has_table for s in deck.slides[0].shapes)
    assert 'Smoke test notes' in deck.slides[0].notes_slide.notes_text_frame.text
    # Resolve every internal relationship as well as parsing all slides.
    for part in deck.part.package.iter_parts():
        for rel in part.rels.values():
            if not rel.is_external:
                assert rel.target_part is not None
    print('Independent python-pptx validation passed:', path)


if __name__ == '__main__':
    generate() if len(sys.argv) == 1 else verify(sys.argv[1])
